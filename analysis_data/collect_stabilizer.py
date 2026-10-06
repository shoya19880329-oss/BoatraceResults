import csv
import time
import threading
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

import openpyxl
import requests

# ============================================================
# 設定
# ============================================================

XLSX = Path(
    "analysis_data/boat_escape_wind_class_20240101_20261005.xlsx"
)

OUTPUT = Path(
    "analysis_data/stabilizer_results.csv"
)

WORKERS = 20
TIMEOUT = 30
MAX_RETRIES = 5

# ============================================================
# 通信用
# ============================================================

thread_local = threading.local()


def get_session():
    if not hasattr(thread_local, "session"):
        session = requests.Session()
        session.headers.update({
            "User-Agent": "Mozilla/5.0"
        })
        thread_local.session = session

    return thread_local.session


# ============================================================
# 1レース取得
# ============================================================

def fetch_race(item):
    hd, jcd, rno = item

    url = (
        "https://www.boatrace.jp/owpc/pc/race/raceresult"
        f"?hd={hd}&jcd={jcd:02d}&rno={rno}"
    )

    session = get_session()

    for attempt in range(1, MAX_RETRIES + 1):

        try:
            response = session.get(
                url,
                timeout=TIMEOUT
            )

            # アクセス過多
            if response.status_code == 429:
                wait = 30 * attempt
                print(
                    f"429: {hd} 場{jcd:02d} {rno}R "
                    f"→ {wait}秒待機"
                )
                time.sleep(wait)
                continue

            response.raise_for_status()

            stable = (
                "有"
                if "安定板使用" in response.text
                else "無"
            )

            return (
                hd,
                jcd,
                rno,
                stable,
                "OK"
            )

        except Exception as e:

            if attempt < MAX_RETRIES:
                wait = 5 * attempt
                time.sleep(wait)

            else:
                return (
                    hd,
                    jcd,
                    rno,
                    "",
                    f"ERROR: {e}"
                )


# ============================================================
# Excelから対象レースを読む
# ============================================================

print("Excelを読み込んでいます...")

wb = openpyxl.load_workbook(
    XLSX,
    read_only=True,
    data_only=True
)

ws = wb["レース明細"]

headers = [
    c.value
    for c in next(
        ws.iter_rows(
            min_row=1,
            max_row=1
        )
    )
]

idx = {
    name: i
    for i, name in enumerate(headers)
}

all_races = []
seen = set()

for row in ws.iter_rows(
    min_row=2,
    values_only=True
):

    d = row[idx["日付"]]
    jcd = int(row[idx["場番号"]])
    rno = int(row[idx["R"]])

    if hasattr(d, "strftime"):
        hd = d.strftime("%Y%m%d")
    else:
        hd = (
            str(d)
            .replace("-", "")
            .replace("/", "")[:8]
        )

    key = (
        hd,
        jcd,
        rno
    )

    if key not in seen:
        seen.add(key)
        all_races.append(key)

wb.close()

print(
    f"対象レース: {len(all_races):,}件"
)

# ============================================================
# 取得済みCSVを読む
# ============================================================

completed = {}

if OUTPUT.exists():

    print(
        "既存の取得結果を読み込んでいます..."
    )

    with OUTPUT.open(
        "r",
        newline="",
        encoding="utf-8-sig"
    ) as f:

        reader = csv.DictReader(f)

        for row in reader:

            if row["状態"] == "OK":

                key = (
                    row["日付"],
                    int(row["場番号"]),
                    int(row["R"])
                )

                completed[key] = row["安定板"]

print(
    f"取得済み: {len(completed):,}件"
)

remaining = [
    race
    for race in all_races
    if race not in completed
]

print(
    f"残り: {len(remaining):,}件"
)

if not remaining:

    print()
    print(
        "全レース取得済みです。"
    )
    raise SystemExit

# ============================================================
# CSV準備
# ============================================================

new_file = not OUTPUT.exists()

csv_file = OUTPUT.open(
    "a",
    newline="",
    encoding="utf-8-sig"
)

writer = csv.writer(csv_file)

if new_file:
    writer.writerow([
        "日付",
        "場番号",
        "R",
        "安定板",
        "状態"
    ])
    csv_file.flush()

# ============================================================
# 本番取得
# ============================================================

print()
print("=" * 60)
print(
    f"{WORKERS}並列で取得開始"
)
print("=" * 60)

start = time.time()

session_completed = 0
session_ok = 0
session_errors = 0
session_stable = 0

with ThreadPoolExecutor(
    max_workers=WORKERS
) as executor:

    futures = {
        executor.submit(
            fetch_race,
            race
        ): race
        for race in remaining
    }

    for future in as_completed(futures):

        result = future.result()

        hd, jcd, rno, stable, status = result

        writer.writerow([
            hd,
            f"{jcd:02d}",
            rno,
            stable,
            status
        ])

        # 1件ごとにディスクへ保存
        csv_file.flush()

        session_completed += 1

        if status == "OK":
            session_ok += 1

            if stable == "有":
                session_stable += 1

        else:
            session_errors += 1

        # 100件ごとに進捗
        if (
            session_completed % 100 == 0
            or session_completed == len(remaining)
        ):

            elapsed = time.time() - start

            speed = (
                session_completed / elapsed
                if elapsed
                else 0
            )

            total_done = (
                len(completed)
                + session_ok
            )

            left = max(
                len(all_races) - total_done,
                0
            )

            eta = (
                left / speed
                if speed
                else 0
            )

            print(
                f"{total_done:,}/{len(all_races):,} "
                f"| 今回 {session_completed:,}件 "
                f"| {speed:.2f}件/秒 "
                f"| 残り約 {eta/3600:.1f}時間 "
                f"| エラー {session_errors}"
            )

csv_file.close()

# ============================================================
# 終了表示
# ============================================================

elapsed = time.time() - start

print()
print("=" * 60)
print("今回の取得終了")
print(
    f"今回処理: {session_completed:,}件"
)
print(
    f"正常: {session_ok:,}件"
)
print(
    f"安定板有: {session_stable:,}件"
)
print(
    f"エラー: {session_errors:,}件"
)
print(
    f"所要時間: {elapsed/3600:.2f}時間"
)
print(
    f"保存先: {OUTPUT}"
)
print("=" * 60)

if session_errors:
    print()
    print(
        "エラー分は次回実行時に再取得されます。"
    )
else:
    print()
    print(
        "エラーはありません。"
    )
