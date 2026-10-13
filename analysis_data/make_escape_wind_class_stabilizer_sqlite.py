import os
from pathlib import Path
from collections import defaultdict
from datetime import date
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# 基本設定
# ============================================================

BASE_DIR = Path(__file__).resolve().parents[1]
DB_FILE = BASE_DIR.parent / "boatrace_master.sqlite"
DATA_DIR = BASE_DIR / "analysis_data"

START_DATE = date(2024, 1, 1)

CLASS_NAMES = {
    1: "A1",
    2: "A2",
    3: "B1",
    4: "B2",
}

STADIUMS = {
    1: "桐生",
    2: "戸田",
    3: "江戸川",
    4: "平和島",
    5: "多摩川",
    6: "浜名湖",
    7: "蒲郡",
    8: "常滑",
    9: "津",
    10: "三国",
    11: "びわこ",
    12: "住之江",
    13: "尼崎",
    14: "鳴門",
    15: "丸亀",
    16: "児島",
    17: "宮島",
    18: "徳山",
    19: "下関",
    20: "若松",
    21: "芦屋",
    22: "福岡",
    23: "唐津",
    24: "大村",
}

WIND_DIRECTIONS = {
    1: "北", 2: "北", 3: "北東",
    4: "東", 5: "東", 6: "東",
    7: "南東",
    8: "南", 9: "南", 10: "南",
    11: "南西",
    12: "西", 13: "西", 14: "西",
    15: "北西", 16: "北",
    17: "無風",
}

WIND_DIRECTIONS_2 = {1: chr(8593), 2: chr(8593)+chr(8599), 3: chr(8599), 4: chr(8594)+chr(8599), 5: chr(8594), 6: chr(8594)+chr(8600), 7: chr(8600), 8: chr(8595)+chr(8600), 9: chr(8595), 10: chr(8595)+chr(8601), 11: chr(8601), 12: chr(8592)+chr(8601), 13: chr(8592), 14: chr(8592)+chr(8598), 15: chr(8598), 16: chr(8593)+chr(8598), 17: chr(28961)+chr(39080)}

TECHNIQUES = {
    1: "逃げ",
    2: "差し",
    3: "まくり",
    4: "まくり差し",
    5: "抜き",
    6: "恵まれ",
}


def wind_speed_band(speed):
    if speed is None:
        return "不明"

    speed = float(speed)

    if 0 <= speed <= 2:
        return "弱"

    if 3 <= speed <= 4:
        return "中"

    if speed >= 5:
        return "強"

    return "不明"


def matches_wind_band(speed, band):
    if speed is None:
        return False

    speed = float(speed)

    if band == "弱":
        return 0 <= speed <= 2

    if band == "中":
        return 3 <= speed <= 4

    if band == "中強":
        return speed >= 3

    if band == "強":
        return speed >= 5

    return False
def open_database():
    if not DB_FILE.exists():
        raise FileNotFoundError(
            f"SQLiteが見つかりません: {DB_FILE}"
        )

    uri = DB_FILE.resolve().as_uri() + "?mode=ro"
    return sqlite3.connect(uri, uri=True)


def get_end_date(conn):
    value = conn.execute(
        "SELECT MAX(race_date) FROM races"
    ).fetchone()[0]

    if not value:
        raise RuntimeError("SQLiteにレースデータがありません")

    return date.fromisoformat(value)


def get_trifecta(conn, race_id):
    rows = conn.execute(
        """
        SELECT combination, amount, popularity
        FROM payouts
        WHERE race_id = ?
          AND bet_type = 'trifecta'
        ORDER BY rowid
        """,
        (race_id,),
    ).fetchall()

    combinations = []
    amounts = []
    popularities = []

    for combination, amount, popularity in rows:
        if combination is not None:
            combinations.append(str(combination))
        if amount is not None:
            amounts.append(str(amount))
        if popularity is not None:
            popularities.append(str(popularity))

    return " / ".join(combinations), " / ".join(amounts), " / ".join(popularities)


def load_detail_extras(conn):
    classes_by_race = defaultdict(dict)

    for race_id, boat_number, class_number in conn.execute(
        """
        SELECT race_id, boat_number, racer_class_number
        FROM race_entries
        ORDER BY race_id, boat_number
        """
    ):
        classes_by_race[race_id][boat_number] = CLASS_NAMES.get(
            class_number,
            "不明",
        )

    exhibition_parts = defaultdict(list)

    for race_id, boat_number in conn.execute(
        """
        SELECT race_id, boat_number
        FROM race_previews
        WHERE exhibition_course_number IS NOT NULL
        ORDER BY race_id, exhibition_course_number, boat_number
        """
    ):
        exhibition_parts[race_id].append(str(boat_number))

    actual_parts = defaultdict(list)

    for race_id, boat_number in conn.execute(
        """
        SELECT race_id, boat_number
        FROM race_entries
        WHERE actual_course_number IS NOT NULL
        ORDER BY race_id, actual_course_number, boat_number
        """
    ):
        actual_parts[race_id].append(str(boat_number))

    exhibition_by_race = {
        race_id: "".join(parts)
        for race_id, parts in exhibition_parts.items()
    }

    actual_by_race = {
        race_id: "".join(parts)
        for race_id, parts in actual_parts.items()
    }

    return classes_by_race, exhibition_by_race, actual_by_race


def collect_races(conn, end_date):
    classes_by_race, exhibition_by_race, actual_by_race = load_detail_extras(conn)
    print()
    print("SQLiteからレース明細を読み込み中...")

    sql = """
    SELECT
        r.race_id,
        r.race_date,
        r.stadium_number,
        r.race_number,
        r.result_wind_direction_number,
        r.result_wind_speed,
        r.technique_number,
        r.stabilizer_used,

        c.boat_number,
        c.racer_number,
        c.racer_name,
        c.racer_class_number,
        c.actual_start_timing,
        c.place_number,

(
    SELECT w.racer_class_number
    FROM race_entries AS w
    WHERE w.race_id = r.race_id
      AND w.place_number = 1
    ORDER BY w.boat_number
    LIMIT 1
) AS winner_class_number

    FROM races AS r

    JOIN race_entries AS c
      ON c.race_id = r.race_id
     AND c.actual_course_number = 1

    WHERE r.race_date BETWEEN ? AND ?

    ORDER BY
        r.race_date,
        r.stadium_number,
        r.race_number
    """

    db_rows = conn.execute(
        sql,
        (
            START_DATE.isoformat(),
            end_date.isoformat(),
        ),
    ).fetchall()

    rows = []

    for index, db_row in enumerate(db_rows, start=1):
        (
            race_id,
            race_date_text,
            stadium_number,
            race_number,
            wind_direction_number,
            wind_speed,
            technique_number,
            stabilizer_used,
            boat_number,
            racer_number,
            racer_name,
            racer_class_number,
            actual_start_timing,
            place_number,
            winner_class_number,
        ) = db_row

        race_date = date.fromisoformat(race_date_text)

        racer_class = CLASS_NAMES.get(
            racer_class_number,
            "不明",
        )

        winner_class = CLASS_NAMES.get(
            winner_class_number,
            "不明",
        )

        technique = TECHNIQUES.get(
            technique_number,
            (
                f"番号{technique_number}"
                if technique_number is not None
                else "不明"
            ),
        )

        is_escape = (
            place_number == 1
            and technique_number == 1
        )

        stabilizer = (
            "あり"
            if stabilizer_used == 1
            else "なし"
        )

        trifecta_combination, trifecta_amount, trifecta_popularity = get_trifecta(
            conn,
            race_id,
        )

        rows.append(
            {
                "日付": race_date,
                "場番号": stadium_number,
                "場": STADIUMS.get(
                    stadium_number,
                    str(stadium_number),
                ),
                "R": race_number,
                "風向番号": wind_direction_number,
                "風向": WIND_DIRECTIONS.get(
                    wind_direction_number,
                    "不明",
                ),
                "風向2": WIND_DIRECTIONS_2.get(wind_direction_number, "不明"),
                "風速": wind_speed,
                "風速帯": wind_speed_band(wind_speed),
                "安定板": stabilizer,
                "1コース艇番": boat_number,
                "1コース選手": racer_name,
                "登録番号": racer_number,
                "級別": racer_class,
                "①級": classes_by_race.get(race_id, {}).get(1, "不明"),
                "②級": classes_by_race.get(race_id, {}).get(2, "不明"),
                "③級": classes_by_race.get(race_id, {}).get(3, "不明"),
                "④級": classes_by_race.get(race_id, {}).get(4, "不明"),
                "⑤級": classes_by_race.get(race_id, {}).get(5, "不明"),
                "⑥級": classes_by_race.get(race_id, {}).get(6, "不明"),
                "展示並び": (
                    "展示なし"
                    if stadium_number == 24
                    and (
                        date(2024, 11, 25) <= race_date <= date(2024, 11, 30)
                        or date(2024, 12, 7) <= race_date <= date(2024, 12, 11)
                    )
                    else exhibition_by_race.get(race_id, "")
                ),
                "本番並び": actual_by_race.get(race_id, ""),
                "ST": actual_start_timing,
                "1コース着順": place_number,
                "決まり手番号": technique_number,
                "決まり手": technique,
                "1着級別": winner_class,
                "3連単組合せ": trifecta_combination,
                "3連単払戻": trifecta_amount,
                "人気": trifecta_popularity,
                "逃げ": 1 if is_escape else 0,
            }
        )

        if index % 10000 == 0:
            print(f"  {index:,}/{len(db_rows):,}レース")

    print(f"対象レース明細: {len(rows):,}件")
    return rows
def make_summary(rows):
    summary = defaultdict(
        lambda: {
            "レース数": 0,
            "逃げ数": 0,
            "3連単払戻合計": 0,
            "3連単件数": 0,
        }
    )

    for row in rows:
        key = (
            row["場番号"],
            row["場"],
            row["風向"],
            row["風向2"],
            row["風速"],
            row["風速帯"],
            row["級別"],
            row["安定板"],
        )

        summary[key]["レース数"] += 1
        summary[key]["逃げ数"] += row["逃げ"]

        payout = row["3連単払戻"]

        if payout:
            payout_parts = [
                x.strip()
                for x in str(payout).split("/")
            ]

            for part in payout_parts:
                cleaned = (
                    part.replace(",", "")
                    .replace("円", "")
                    .strip()
                )

                if cleaned.isdigit():
                    summary[key]["3連単払戻合計"] += int(cleaned)
                    summary[key]["3連単件数"] += 1

    result = []

    for key, values in summary.items():
        (
            stadium_number,
            stadium,
            wind_direction,
            wind_direction2,
            wind_speed,
            wind_band,
            racer_class,
            stabilizer,
        ) = key

        races = values["レース数"]
        escapes = values["逃げ数"]
        trifecta_count = values["3連単件数"]

        avg_trifecta = (
            values["3連単払戻合計"] / trifecta_count
            if trifecta_count
            else 0
        )

        result.append(
            {
                "場番号": stadium_number,
                "場": stadium,
                "風向": wind_direction,
                "風向2": wind_direction2,
                "風速": wind_speed,
                "風速帯": wind_band,
                "級別": racer_class,
                "安定板": stabilizer,
                "レース数": races,
                "逃げ数": escapes,
                "逃げ率": escapes / races if races else 0,
                "平均3連単払戻": avg_trifecta,
            }
        )

    result.sort(
        key=lambda x: (
            x["場番号"] if x["場番号"] is not None else 999,
            x["風向"],
            x["風速"] if x["風速"] is not None else 999,
            x["級別"],
            x["安定板"],
        )
    )

    return result


def make_class_summary(rows):
    summary = defaultdict(
        lambda: {
            "レース数": 0,
            "逃げ数": 0,
        }
    )

    for row in rows:
        key = (
            row["場番号"],
            row["場"],
            row["級別"],
            row["安定板"],
        )

        summary[key]["レース数"] += 1
        summary[key]["逃げ数"] += row["逃げ"]

    result = []

    for key, values in summary.items():
        stadium_number, stadium, racer_class, stabilizer = key

        races = values["レース数"]
        escapes = values["逃げ数"]

        result.append(
            {
                "場番号": stadium_number,
                "場": stadium,
                "級別": racer_class,
                "安定板": stabilizer,
                "レース数": races,
                "逃げ数": escapes,
                "逃げ率": escapes / races if races else 0,
            }
        )

    result.sort(
        key=lambda x: (
            x["場番号"] if x["場番号"] is not None else 999,
            x["級別"],
            x["安定板"],
        )
    )

    return result
def make_wind_band_summary(rows):
    classes = ["A1", "A2", "B1", "B2"]
    bands = ["弱", "中", "中強", "強"]
    stabilizers = ["なし", "あり"]

    summary = {}

    for stadium_number in range(1, 25):
        for racer_class in classes:
            for band in bands:
                for stabilizer in stabilizers:
                    summary[
                        (
                            stadium_number,
                            racer_class,
                            band,
                            stabilizer,
                        )
                    ] = {
                        "レース数": 0,
                        "逃げ数": 0,
                    }

    for row in rows:
        stadium_number = row["場番号"]
        racer_class = row["級別"]
        speed = row["風速"]
        stabilizer = row["安定板"]

        if racer_class not in classes:
            continue

        if stabilizer not in stabilizers:
            continue

        for band in bands:
            if matches_wind_band(speed, band):
                key = (
                    stadium_number,
                    racer_class,
                    band,
                    stabilizer,
                )

                if key not in summary:
                    continue

                summary[key]["レース数"] += 1
                summary[key]["逃げ数"] += row["逃げ"]

    return summary


def style_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze

    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.row_dimensions[1].height = 24


def center_all_cells(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )


def set_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def format_rate(races, escapes):
    if races:
        return f"{escapes / races * 100:.1f}%({races})"

    return "0.0%(0)"
def create_workbook(
    rows,
    summary_rows,
    class_summary_rows,
    wind_band_summary,
    end_date,
):
    print()
    print("Excel作成開始...")

    auto_output = os.environ.get("BR_EXCEL_OUTPUT")

    if auto_output:
        output_file = Path(auto_output)
    else:
        output_file = DATA_DIR / (
            "boat_escape_wind_class_stabilizer_"
            f"{START_DATE:%Y%m%d}_{end_date:%Y%m%d}.xlsx"
        )

    wb = Workbook()

    # ========================================================
    # 1. レース明細
    # ========================================================

    ws = wb.active
    ws.title = "レース明細"

    detail_headers = [
        "日付",
        "場番号",
        "場",
        "R",
        "風向",
        "風向2",
        "風速",
        "風速帯",
        "安定板",
        "1コース艇番",
        "1コース選手",
        "登録番号",
        "級別",
        "①級",
        "②級",
        "③級",
        "④級",
        "⑤級",
        "⑥級",
        "展示並び",
        "本番並び",
        "ST",
        "1コース着順",
        "決まり手番号",
        "決まり手",
        "1着級別",
        "3連単組合せ",
        "3連単払戻",
        "人気",
        "逃げ",
    ]

    ws.append(detail_headers)

    for row in rows:
        values = []

        for header in detail_headers:
            value = row[header]

            if header == "風速" and value is not None:
                value = f"{value:g}m"

            values.append(value)

        ws.append(values)

    for cell in ws["A"][1:]:
        cell.number_format = "yyyy/mm/dd"

    for cell in ws["M"][1:]:
        cell.number_format = "0.00"

    style_sheet(ws)

    set_widths(
        ws,
        {
            "A": 12,
            "B": 8,
            "C": 10,
            "D": 6,
            "E": 10,
            "F": 8,
            "G": 10,
            "H": 8,
            "I": 12,
            "J": 18,
            "K": 12,
            "L": 8,
            "M": 8,
            "N": 12,
            "O": 12,
            "P": 14,
            "Q": 10,
            "R": 16,
            "S": 14,
            "T": 8,
        },
    )

    print(f"  レース明細: {len(rows):,}行")

    # ========================================================
    # 2. 風向風速級別集計
    # ========================================================

    ws2 = wb.create_sheet("風向風速級別集計")

    summary_headers = [
        "場番号",
        "場",
        "風向",
        "風向2",
        "風速",
        "風速帯",
        "級別",
        "安定板",
        "レース数",
        "逃げ数",
        "逃げ率",
        "平均3連単払戻",
    ]

    ws2.append(summary_headers)

    for row in summary_rows:
        values = [
            row[header]
            for header in summary_headers
        ]

        if values[4] is not None:
            values[4] = f"{values[4]:g}m"

        ws2.append(values)

    for cell in ws2["K"][1:]:
        cell.number_format = "0.0%"

    for cell in ws2["L"][1:]:
        cell.number_format = '#,##0"円"'

    style_sheet(ws2)

    set_widths(
        ws2,
        {
            "A": 8,
            "B": 10,
            "C": 12,
            "D": 8,
            "E": 10,
            "F": 8,
            "G": 10,
            "H": 12,
            "I": 10,
            "J": 12,
            "K": 16,
        },
    )

    print(f"  風向風速級別集計: {len(summary_rows):,}行")

    # ========================================================
    # 3. 場級別集計
    # ========================================================

    ws3 = wb.create_sheet("場級別集計")

    class_headers = [
        "場番号",
        "場",
        "級別",
        "安定板",
        "レース数",
        "逃げ数",
        "逃げ率",
    ]

    ws3.append(class_headers)

    for row in class_summary_rows:
        ws3.append(
            [
                row[header]
                for header in class_headers
            ]
        )

    for cell in ws3["G"][1:]:
        cell.number_format = "0.0%"

    style_sheet(ws3)

    set_widths(
        ws3,
        {
            "A": 8,
            "B": 10,
            "C": 8,
            "D": 10,
            "E": 12,
            "F": 10,
            "G": 12,
        },
    )

    print(f"  場級別集計: {len(class_summary_rows):,}行")
    # ========================================================
    # 4. 風速帯別集計
    # ========================================================

    ws4 = wb.create_sheet("風速帯別集計")

    classes = ["A1", "A2", "B1", "B2"]
    bands = ["弱", "中", "中強", "強"]
    stabilizers = ["なし", "あり"]

    headers = ["場"]

    for racer_class in classes:
        for band in bands:
            for stabilizer in stabilizers:
                headers.append(
                    f"{racer_class}({band})-{stabilizer}"
                )

    ws4.append(headers)

    for stadium_number in range(1, 25):
        values = [
            STADIUMS[stadium_number]
        ]

        for racer_class in classes:
            for band in bands:
                for stabilizer in stabilizers:
                    data = wind_band_summary[
                        (
                            stadium_number,
                            racer_class,
                            band,
                            stabilizer,
                        )
                    ]

                    values.append(
                        format_rate(
                            data["レース数"],
                            data["逃げ数"],
                        )
                    )

        ws4.append(values)

    style_sheet(ws4)

    ws4.column_dimensions["A"].width = 12

    for col in range(2, len(headers) + 1):
        ws4.column_dimensions[
            get_column_letter(col)
        ].width = 17

    print(
        "  風速帯別集計: "
        "24場 × 4級別 × 4風速帯 × 2安定板条件"
    )
    # ========================================================
    # 5. 人気
    # ========================================================

    ws5 = wb.create_sheet("人気")

    popularity_headers = [
        "場・風向",
        "人気",
        "3連単払戻",
        "風速2ｍ",
        "風速帯中",
        "風速帯強",
        "人気(安)",
        "3連単払戻(安)",
        "風速2ｍ(安)",
        "風速帯中(安)",
        "風速帯強(安)",
    ]
    ws5.append(popularity_headers)

    popularity_directions = [
        "東", "西", "南", "北",
        "南東", "北東", "南西", "北西",
    ]

    def numeric_parts(value):
        if value is None or value == "":
            return []
        result = []
        for part in str(value).split("/"):
            cleaned = (
                part.replace(",", "")
                .replace("円", "")
                .strip()
            )
            try:
                result.append(float(cleaned))
            except ValueError:
                pass
        return result

    for stadium_number in range(1, 25):
        stadium = STADIUMS[stadium_number]

        for direction in popularity_directions:
            matched = [
                row for row in rows
                if row["場番号"] == stadium_number
                and row["風向"] == direction
            ]

            output = [f"{stadium}({direction})"]

            for stabilizer in ["なし", "あり"]:
                target = [
                    row for row in matched
                    if row["安定板"] == stabilizer
                ]

                speed_2_or_more = [
                    row for row in target
                    if row["風速"] is not None
                    and float(row["風速"]) >= 2
                ]

                popularity_values = []
                payout_values = []

                for row in speed_2_or_more:
                    popularity_values.extend(
                        numeric_parts(row["人気"])
                    )
                    payout_values.extend(
                        numeric_parts(row["3連単払戻"])
                    )

                avg_popularity = (
                    sum(popularity_values) / len(popularity_values)
                    if popularity_values else None
                )
                avg_payout = (
                    sum(payout_values) / len(payout_values)
                    if payout_values else None
                )

                speed_2_count = sum(
                    1 for row in target
                    if row["風速"] is not None
                    and float(row["風速"]) == 2
                )
                medium_count = sum(
                    1 for row in target
                    if row["風速帯"] == "中"
                )
                strong_count = sum(
                    1 for row in target
                    if row["風速帯"] == "強"
                )

                output.extend([
                    avg_popularity,
                    avg_payout,
                    speed_2_count,
                    medium_count,
                    strong_count,
                ])

            ws5.append(output)

    style_sheet(ws5)
    ws5.column_dimensions["A"].width = 18

    for col in range(2, len(popularity_headers) + 1):
        ws5.column_dimensions[get_column_letter(col)].width = 18

    for row_cells in ws5.iter_rows(
        min_row=2,
        min_col=2,
        max_col=2,
    ):
        row_cells[0].number_format = "0.0"

    for row_cells in ws5.iter_rows(
        min_row=2,
        min_col=7,
        max_col=7,
    ):
        row_cells[0].number_format = "0.0"

    for row_cells in ws5.iter_rows(
        min_row=2,
        min_col=3,
        max_col=3,
    ):
        row_cells[0].number_format = "#,##0"

    for row_cells in ws5.iter_rows(
        min_row=2,
        min_col=8,
        max_col=8,
    ):
        row_cells[0].number_format = "#,##0"

    print("  人気: 24場 × 8風向")

    # ========================================================
    # 6. 説明
    # ========================================================

    ws6 = wb.create_sheet("説明")

    escape_count = sum(
        row["逃げ"]
        for row in rows
    )

    stabilizer_yes = sum(
        1
        for row in rows
        if row["安定板"] == "あり"
    )

    stabilizer_no = sum(
        1
        for row in rows
        if row["安定板"] == "なし"
    )

    unknown_grade = sum(
        1
        for row in rows
        if row["級別"] == "不明"
    )

    unknown_winner_grade = sum(
        1
        for row in rows
        if row["1着級別"] == "不明"
    )

    info_rows = [
        ["項目", "内容"],
        ["指定開始日", START_DATE],
        ["実データ終了日", end_date],
        ["分析対象レース数", len(rows)],
        ["逃げ数", escape_count],
        [
            "全体逃げ率",
            escape_count / len(rows) if rows else 0,
        ],
        ["安定板なしレース数", stabilizer_no],
        ["安定板ありレース数", stabilizer_yes],
        ["1コース級別不明件数", unknown_grade],
        ["1着級別不明件数", unknown_winner_grade],
        [
            "1コースの定義",
            "actual_course_number = 1 の実際の1コース進入選手",
        ],
        [
            "逃げの定義",
            "実際の1コース選手が1着、かつ決まり手番号=1（逃げ）",
        ],
        [
            "級別",
            "SQLiteに保存された各レースの公式番組級別を使用",
        ],
        [
            "安定板",
            "SQLiteの stabilizer_used を使用（0=なし、1=あり）",
        ],
        ["風速帯：弱", "0m～2m"],
        ["風速帯：中", "3m～4m"],
        [
            "風速帯：中強",
            "3m以上（中・強と重複して集計）",
        ],
        ["風速帯：強", "5m以上"],
        [
            "風向風速級別集計",
            "場 × 風向 × 風速 × 1コース選手級別 × 安定板で集計",
        ],
        [
            "場級別集計",
            "場 × 1コース選手級別 × 安定板で集計",
        ],
        [
            "風速帯別集計",
            "場 × 1コース選手級別 × 風速帯 × 安定板で逃げ率を表示",
        ],
        [
            "風速帯別集計の例",
            "62.0%(100) = 対象100レース中、逃げ62レース",
        ],
    ]

    for info in info_rows:
        ws6.append(info)

    for row_number in (2, 3):
        ws6.cell(
            row=row_number,
            column=2,
        ).number_format = "yyyy/mm/dd"

    ws5["B6"].number_format = "0.0%"

    style_sheet(ws5)

    ws6.column_dimensions["A"].width = 30
    ws6.column_dimensions["B"].width = 80

    # ========================================================
    # 全シート中央揃え
    # ========================================================

    for sheet in wb.worksheets:
        center_all_cells(sheet)

    # ========================================================
    # 保存
    # ========================================================

    wb.save(output_file)

    print()
    print("=" * 60)
    print("Excel完成")
    print(f"保存先: {output_file}")
    print(
        f"ファイルサイズ: "
        f"{output_file.stat().st_size:,} bytes"
    )
    print("=" * 60)

    return output_file
def main():
    print("=" * 60)
    print("BOAT RACE 1コース逃げ率・安定板分析")
    print(f"SQLite: {DB_FILE}")
    print("=" * 60)

    conn = open_database()

    try:
        end_date = get_end_date(conn)

        print(
            f"分析期間: "
            f"{START_DATE} ～ {end_date}"
        )

        rows = collect_races(
            conn,
            end_date,
        )

        if not rows:
            raise RuntimeError(
                "分析対象レースがありません"
            )

        print()
        print("集計中...")

        summary_rows = make_summary(rows)

        class_summary_rows = make_class_summary(
            rows
        )

        wind_band_summary = make_wind_band_summary(
            rows
        )

        create_workbook(
            rows,
            summary_rows,
            class_summary_rows,
            wind_band_summary,
            end_date,
        )

    finally:
        conn.close()


if __name__ == "__main__":
    main()
