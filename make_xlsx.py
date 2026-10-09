import csv
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

with open("boat_racer_complete.csv", encoding="utf-8") as f:
    rows = list(csv.reader(f))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "2コース選手分析"

for i, row in enumerate(rows):
    if i == 0:
        ws.append(row)
    else:
        converted = []
        for j, v in enumerate(row):
            if j == 1:
                converted.append(v)
            elif j == 0:
                converted.append(int(v))
            else:
                try:
                    converted.append(float(v) if "." in v else int(v))
                except:
                    converted.append(v)
        ws.append(converted)

ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions

for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

widths = {
    "A":10,"B":18,"C":16,"D":16,"E":12,
    "F":14,"G":14,"H":14,"I":14,"J":12,
    "K":12,"L":12,"M":12,"N":12,"O":16,
    "P":16,"Q":12,"R":12,"S":12,"T":12
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(horizontal="center")

wb.save("boat_racer_complete.xlsx")
print("Excel完成:", ws.max_row-1, "選手 /", ws.max_column, "項目")
